from __future__ import annotations

import json
import re
import sys
import traceback
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Callable, Iterable

import fitz
import numpy as np
from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PySide6.QtCore import Qt, QThread, Signal
from PySide6.QtGui import QImage, QPixmap
from PySide6.QtWidgets import (
    QApplication, QCheckBox, QFileDialog, QHBoxLayout, QLabel, QMainWindow,
    QMessageBox, QPushButton, QProgressBar, QSplitter, QTableWidget,
    QTableWidgetItem, QVBoxLayout, QWidget
)

APP_NAME = "3년차 하자접수 PDF → Excel"
APP_VERSION = "0.1.0"

HEADER_ROIS = {
    "dong_ho": (0.040, 0.085, 0.280, 0.125),
    "owner": (0.275, 0.085, 0.470, 0.125),
    "phone": (0.465, 0.085, 0.720, 0.125),
    "weekend": (0.715, 0.085, 0.945, 0.125),
}

PAGE1_ROWS = [
    ("냉방·난방·환기·공기조화설비공사", "자체보일러·에어컨·보일러 등", 0.355, 0.402),
    ("급배수 및 위생설비공사", "온수설비", 0.403, 0.440),
    ("급배수 및 위생설비공사", "난방설비", 0.440, 0.475),
    ("급배수 및 위생설비공사", "세대 내 등", 0.475, 0.515),
    ("가스설비공사", "가스설비 등", 0.515, 0.557),
    ("창호공사", "PL창호·복층유리·창호유리·방화문 등", 0.557, 0.617),
    ("정보통신공사", "TV공청설비·통신설비·자동소화기 등", 0.617, 0.685),
    ("홈네트워크공사", "홈네트워크기기·단지공용시스템 등", 0.685, 0.745),
    ("단열공사", "벽체 및 천장 단열공사 등", 0.745, 0.802),
    ("콘크리트공사", "벽체 및 천장 단열공사 등", 0.802, 0.855),
]

PAGE2_ROWS = [
    ("목공사", "구조체·수장목공사 등", 0.140, 0.222),
    ("전기 및 전력설비공사", "배관배선·수변전반·전기기기 등", 0.222, 0.305),
    ("잡공사", "우편함·무인택배·몰딩류 등", 0.305, 0.375),
    ("소방시설공사", "소화설비·제연설비·방재설비·피난기구·감지기 등", 0.375, 0.515),
    ("방수공사", "천장", 0.515, 0.570),
    ("방수공사", "벽", 0.570, 0.625),
    ("방수공사", "바닥", 0.625, 0.682),
    ("기타", "기타", 0.682, 0.785),
]

LOCATIONS = [
    "거실", "주방", "침실1", "침실2", "침실3", "안방", "드레스룸", "현관",
    "실외기실", "대피공간", "공용욕실", "부부욕실", "욕실1", "욕실2", "발코니",
    "앞발코니", "뒷발코니", "다용도실", "세탁실", "팬트리", "복도", "전실", "베란다"
]
PARTS = [
    "벽", "바닥", "천장", "벽부등", "배전함", "보일러", "분배기", "신발장", "욕실",
    "창틀", "창문", "창호", "문", "문틀", "수전", "세면대", "양변기", "변기", "배수구",
    "트랩", "배관", "콘센트", "스위치", "조명", "등기구", "감지기", "환풍기", "후드",
    "도어락", "인터폰", "싱크대", "상판", "타일", "마루", "도배", "몰딩", "실리콘",
    "방화문", "유리", "난간", "레일", "손잡이"
]
DEFECT_TYPES = [
    "깨짐", "파손", "들뜸", "오염", "작동불량", "누수", "균열", "크랙", "소음", "변색",
    "탈락", "마감불량", "처짐", "흔들림", "벌어짐", "틈", "들림", "찍힘", "스크래치",
    "막힘", "역류", "결로", "곰팡이", "부식", "녹", "미작동", "점등불량", "개폐불량",
    "수압불량", "배수불량", "수평불량", "단차", "유격", "미시공", "누락"
]
ALIASES = {
    "안방": "침실1", "부부화장실": "부부욕실", "안방화장실": "부부욕실",
    "공용화장실": "공용욕실", "거실화장실": "공용욕실",
    "화장실1": "공용욕실", "화장실2": "부부욕실",
}


@dataclass
class DefectRow:
    dong: str = ""
    ho: str = ""
    owner: str = ""
    phone: str = ""
    weekend: str = ""
    work: str = ""
    subwork: str = ""
    defect: str = ""
    location: str = ""
    part: str = ""
    defect_type: str = ""
    note: str = ""
    source_page: str = ""
    confidence: float = 0.0


class OCREngine:
    def __init__(self):
        self._ocr = None

    def _load(self):
        if self._ocr is not None:
            return
        from paddleocr import PaddleOCR
        self._ocr = PaddleOCR(
            lang="korean",
            use_doc_orientation_classify=False,
            use_doc_unwarping=False,
            use_textline_orientation=False,
        )

    def read_lines(self, image: Image.Image) -> list[tuple[str, float]]:
        self._load()
        arr = np.array(image.convert("RGB"))
        try:
            results = self._ocr.predict(arr)
            lines: list[tuple[str, float]] = []
            for r in results:
                data = getattr(r, "json", None)
                if callable(data):
                    data = data()
                if isinstance(data, str):
                    data = json.loads(data)
                if isinstance(data, dict):
                    root = data.get("res", data)
                    texts = root.get("rec_texts", [])
                    scores = root.get("rec_scores", [])
                    for i, txt in enumerate(texts):
                        txt = clean_text(str(txt))
                        if txt:
                            score = float(scores[i]) if i < len(scores) else 0.0
                            lines.append((txt, score))
            if lines:
                return lines
        except Exception:
            pass

        results = self._ocr.ocr(arr, cls=False)
        lines = []
        if not results:
            return lines
        page = results[0] if isinstance(results, list) and results else results
        if page:
            for item in page:
                try:
                    txt, score = item[1][0], item[1][1]
                    txt = clean_text(str(txt))
                    if txt:
                        lines.append((txt, float(score)))
                except Exception:
                    continue
        return lines


def clean_text(s: str) -> str:
    s = s.replace("\u3000", " ")
    s = re.sub(r"\s+", " ", s)
    return s.strip(" |,.;:_-")


def crop_norm(img: Image.Image, roi: tuple[float, float, float, float]) -> Image.Image:
    w, h = img.size
    x1, y1, x2, y2 = roi
    return img.crop((int(w*x1), int(h*y1), int(w*x2), int(h*y2)))


def upscale(img: Image.Image, factor: float = 2.0) -> Image.Image:
    return img.resize((int(img.width*factor), int(img.height*factor)), Image.Resampling.LANCZOS)


def pdf_pages(pdf_path: str, dpi: int = 220) -> list[Image.Image]:
    doc = fitz.open(pdf_path)
    pages = []
    zoom = dpi / 72.0
    mat = fitz.Matrix(zoom, zoom)
    for p in doc:
        pix = p.get_pixmap(matrix=mat, alpha=False)
        pages.append(Image.frombytes("RGB", [pix.width, pix.height], pix.samples))
    return pages


def best_text(lines: list[tuple[str, float]]) -> tuple[str, float]:
    if not lines:
        return "", 0.0
    return clean_text(" ".join(t for t, _ in lines)), sum(s for _, s in lines) / len(lines)


def parse_dong_ho(text: str) -> tuple[str, str]:
    text = text.replace("O", "0").replace("o", "0")
    m = re.search(r"(\d{3,4})\s*동?.*?(\d{3,4})\s*호?", text)
    if m:
        return m.group(1), m.group(2)
    nums = re.findall(r"\d{2,4}", text)
    if len(nums) >= 2:
        return nums[0].lstrip("0") or "0", nums[1].lstrip("0") or "0"
    return (nums[0], "") if nums else ("", "")


def normalize_phone(text: str) -> str:
    digits = re.sub(r"\D", "", text)
    if len(digits) == 11 and digits.startswith("010"):
        return f"{digits[:3]}-{digits[3:7]}-{digits[7:]}"
    if len(digits) == 10 and digits.startswith("01"):
        return f"{digits[:3]}-{digits[3:6]}-{digits[6:]}"
    return text


def parse_weekend(text: str) -> str:
    t = text.replace(" ", "")
    if "불가능" in t or "불가" in t or "×" in t:
        return "불가능"
    if "가능" in t or "○" in t or "✓" in t or "V" in t:
        return "가능"
    return "확인필요" if t else ""


def normalize_location(text: str) -> str:
    for k, v in ALIASES.items():
        if k in text:
            return v
    for loc in LOCATIONS:
        if loc in text:
            return loc
    m = re.search(r"침실\s*([123])", text)
    return f"침실{m.group(1)}" if m else ""


def find_keyword(text: str, words: Iterable[str]) -> str:
    for w in words:
        if w in text:
            return w
    return ""


def split_defect_lines(lines: list[tuple[str, float]]) -> list[tuple[str, float]]:
    out = []
    for txt, score in lines:
        parts = [clean_text(x) for x in re.split(r"[•●▶▷;]|\s{3,}", txt) if clean_text(x)]
        for p in parts:
            if len(p) >= 2:
                out.append((p, score))
    return out


def parse_form_pair(page1, page2, page1_no, page2_no, ocr, progress):
    header, header_conf = {}, {}
    for name, roi in HEADER_ROIS.items():
        progress(f"기본정보 인식: {name}")
        txt, conf = best_text(ocr.read_lines(upscale(crop_norm(page1, roi), 2.4)))
        header[name], header_conf[name] = txt, conf

    dong, ho = parse_dong_ho(header.get("dong_ho", ""))
    owner = re.sub(r"\([^)]*\)", "", header.get("owner", "")).strip()
    phone = normalize_phone(header.get("phone", ""))
    weekend = parse_weekend(header.get("weekend", ""))
    rows = []

    def process_rows(img, specs, page_no):
        for work, subwork, y1, y2 in specs:
            progress(f"{page_no}페이지: {work} / {subwork}")
            lines = ocr.read_lines(upscale(crop_norm(img, (0.365, y1, 0.955, y2)), 2.2))
            for txt, conf in split_defect_lines(lines):
                if len(txt) < 2 or txt in {work, subwork}:
                    continue
                location = normalize_location(txt)
                part = find_keyword(txt, PARTS)
                dtype = find_keyword(txt, DEFECT_TYPES)
                notes = []
                if conf < 0.72: notes.append("수기 판독 신뢰도 낮음 - 원본 확인 필요")
                if not location: notes.append("위치 확인 필요")
                if not part: notes.append("부위 확인 필요")
                if not dtype: notes.append("하자유형 확인 필요")
                if not dong or not ho: notes.append("동·호수 확인 필요")
                if header_conf.get("owner", 0) < 0.70: notes.append("소유자성명 확인 필요")
                if header_conf.get("phone", 0) < 0.75: notes.append("연락처 확인 필요")
                rows.append(DefectRow(
                    dong=dong, ho=ho, owner=owner, phone=phone, weekend=weekend,
                    work=work, subwork=subwork, defect=txt, location=location, part=part,
                    defect_type=dtype, note=" / ".join(dict.fromkeys(notes)),
                    source_page=str(page_no), confidence=round(conf, 3)
                ))

    process_rows(page1, PAGE1_ROWS, page1_no)
    if page2 is not None and page2_no is not None:
        process_rows(page2, PAGE2_ROWS, page2_no)

    if not rows:
        rows.append(DefectRow(
            dong=dong, ho=ho, owner=owner, phone=phone, weekend=weekend,
            note="하자내용 자동 인식 없음 - 원본 확인 필요",
            source_page=f"{page1_no}" + (f",{page2_no}" if page2_no else ""), confidence=0.0
        ))
    return rows


class AnalyzeThread(QThread):
    progress = Signal(int, int, str)
    finished_ok = Signal(object, object)
    failed = Signal(str)

    def __init__(self, pdf_path: str):
        super().__init__()
        self.pdf_path = pdf_path

    def run(self):
        try:
            pages = pdf_pages(self.pdf_path)
            ocr = OCREngine()
            rows = []
            total_pairs = (len(pages) + 1) // 2
            for pair_idx in range(total_pairs):
                p1_idx = pair_idx * 2
                p2_idx = p1_idx + 1
                p1 = pages[p1_idx]
                p2 = pages[p2_idx] if p2_idx < len(pages) else None
                def cb(msg):
                    self.progress.emit(pair_idx + 1, total_pairs, msg)
                rows.extend(parse_form_pair(p1, p2, p1_idx + 1, p2_idx + 1 if p2 else None, ocr, cb))
            self.finished_ok.emit(rows, pages)
        except Exception:
            self.failed.emit(traceback.format_exc())


class MainWindow(QMainWindow):
    COLS = [
        ("동", "dong"), ("호수", "ho"), ("소유자성명", "owner"), ("연락처", "phone"),
        ("주말작업가능여부", "weekend"), ("3년차 공사명", "work"), ("세부공종", "subwork"),
        ("하자내용", "defect"), ("위치", "location"), ("부위", "part"),
        ("하자유형", "defect_type"), ("비고", "note"), ("원본페이지", "source_page"),
        ("OCR신뢰도", "confidence")
    ]

    def __init__(self):
        super().__init__()
        self.setWindowTitle(f"{APP_NAME} v{APP_VERSION}")
        self.resize(1500, 900)
        self.pdf_path, self.rows, self.pages, self.current_page = "", [], [], 0
        self.thread = None

        root = QWidget(); self.setCentralWidget(root); layout = QVBoxLayout(root)
        top = QHBoxLayout()
        self.open_btn = QPushButton("PDF 선택")
        self.run_btn = QPushButton("자동 인식")
        self.export_btn = QPushButton("Excel 저장"); self.export_btn.setEnabled(False)
        self.prev_btn = QPushButton("◀ 이전 페이지"); self.prev_btn.setEnabled(False)
        self.next_btn = QPushButton("다음 페이지 ▶"); self.next_btn.setEnabled(False)
        self.low_only = QCheckBox("확인 필요 행만 보기")
        for w in [self.open_btn, self.run_btn, self.export_btn, self.prev_btn, self.next_btn, self.low_only]:
            top.addWidget(w)
        top.addStretch(1); layout.addLayout(top)

        self.status = QLabel("PDF를 선택해 주세요.")
        self.progress = QProgressBar(); self.progress.setRange(0, 100)
        layout.addWidget(self.status); layout.addWidget(self.progress)

        splitter = QSplitter(Qt.Horizontal)
        self.preview = QLabel("원본 미리보기"); self.preview.setAlignment(Qt.AlignCenter)
        self.preview.setMinimumWidth(420)
        self.preview.setStyleSheet("QLabel{background:#efefef;border:1px solid #bbb;}")
        splitter.addWidget(self.preview)
        self.table = QTableWidget(0, len(self.COLS))
        self.table.setHorizontalHeaderLabels([x[0] for x in self.COLS])
        self.table.setAlternatingRowColors(True); self.table.setWordWrap(True)
        self.table.verticalHeader().setDefaultSectionSize(46)
        splitter.addWidget(self.table)
        splitter.setStretchFactor(0, 1); splitter.setStretchFactor(1, 3)
        layout.addWidget(splitter, 1)

        self.open_btn.clicked.connect(self.choose_pdf)
        self.run_btn.clicked.connect(self.analyze)
        self.export_btn.clicked.connect(self.export_excel)
        self.prev_btn.clicked.connect(lambda: self.change_page(-1))
        self.next_btn.clicked.connect(lambda: self.change_page(1))
        self.low_only.toggled.connect(self.apply_filter)
        self.table.cellClicked.connect(self.on_cell_clicked)

    def choose_pdf(self):
        path, _ = QFileDialog.getOpenFileName(self, "하자접수 PDF 선택", "", "PDF (*.pdf)")
        if not path: return
        self.pdf_path = path
        self.status.setText(f"선택: {path}"); self.progress.setValue(0)
        self.rows = []; self.pages = []; self.table.setRowCount(0)
        self.export_btn.setEnabled(False)

    def analyze(self):
        if not self.pdf_path:
            self.choose_pdf()
            if not self.pdf_path: return
        self.run_btn.setEnabled(False); self.open_btn.setEnabled(False)
        self.status.setText("OCR 모델 준비 중... 최초 실행은 모델 다운로드로 시간이 걸릴 수 있습니다.")
        self.thread = AnalyzeThread(self.pdf_path)
        self.thread.progress.connect(self.on_progress)
        self.thread.finished_ok.connect(self.on_done)
        self.thread.failed.connect(self.on_failed)
        self.thread.start()

    def on_progress(self, current, total, msg):
        self.progress.setValue(int(current / max(total, 1) * 100))
        self.status.setText(f"세대 {current}/{total} - {msg}")

    def on_done(self, rows, pages):
        self.rows, self.pages = rows, pages
        self.progress.setValue(100)
        self.status.setText(f"완료: {len(rows)}개 행 인식. 노란색 행은 원본 확인을 권장합니다.")
        self.run_btn.setEnabled(True); self.open_btn.setEnabled(True)
        self.export_btn.setEnabled(True)
        self.prev_btn.setEnabled(bool(pages)); self.next_btn.setEnabled(bool(pages))
        self.populate_table()
        if pages: self.current_page = 0; self.show_page(0)

    def on_failed(self, tb):
        self.run_btn.setEnabled(True); self.open_btn.setEnabled(True)
        self.status.setText("인식 중 오류가 발생했습니다.")
        QMessageBox.critical(self, "오류", tb[-5000:])

    def populate_table(self):
        self.table.setRowCount(len(self.rows))
        for r, row in enumerate(self.rows):
            d = asdict(row)
            needs_check = bool(row.note) or row.confidence < 0.72
            for c, (_, key) in enumerate(self.COLS):
                val = d[key]
                if key == "confidence": val = f"{float(val):.3f}"
                item = QTableWidgetItem(str(val))
                if key in {"source_page", "confidence"}: item.setTextAlignment(Qt.AlignCenter)
                if needs_check: item.setBackground(Qt.GlobalColor.yellow)
                self.table.setItem(r, c, item)
        self.table.resizeColumnsToContents()
        for c in range(self.table.columnCount()):
            if self.table.columnWidth(c) > 420: self.table.setColumnWidth(c, 420)
        self.apply_filter()

    def apply_filter(self):
        only = self.low_only.isChecked()
        for r, row in enumerate(self.rows):
            self.table.setRowHidden(r, only and not (row.note or row.confidence < 0.72))

    def sync_from_table(self):
        for r, row in enumerate(self.rows):
            for c, (_, key) in enumerate(self.COLS):
                item = self.table.item(r, c)
                if item is None: continue
                text = item.text().strip()
                if key == "confidence":
                    try: setattr(row, key, float(text))
                    except ValueError: pass
                else:
                    setattr(row, key, text)

    def on_cell_clicked(self, row, col):
        try:
            page = int(self.table.item(row, 12).text().split(",")[0]) - 1
            if 0 <= page < len(self.pages):
                self.current_page = page; self.show_page(page)
        except Exception:
            pass

    def change_page(self, delta):
        if not self.pages: return
        self.current_page = max(0, min(len(self.pages)-1, self.current_page + delta))
        self.show_page(self.current_page)

    def show_page(self, idx):
        if not self.pages: return
        img = self.pages[idx]
        max_w = max(400, self.preview.width()-20); max_h = max(500, self.preview.height()-20)
        copy = img.copy(); copy.thumbnail((max_w, max_h), Image.Resampling.LANCZOS)
        data = copy.tobytes("raw", "RGB")
        qimg = QImage(data, copy.width, copy.height, copy.width*3, QImage.Format_RGB888).copy()
        self.preview.setPixmap(QPixmap.fromImage(qimg))
        self.status.setText(self.status.text().split(" | 페이지")[0] + f" | 페이지 {idx+1}/{len(self.pages)}")

    def export_excel(self):
        self.sync_from_table()
        default = Path(self.pdf_path).with_suffix("").name + "_정리.xlsx"
        out, _ = QFileDialog.getSaveFileName(self, "Excel 저장", default, "Excel (*.xlsx)")
        if not out: return
        if not out.lower().endswith(".xlsx"): out += ".xlsx"
        wb = Workbook(); ws = wb.active; ws.title = "3년차 하자접수"
        headers = [x[0] for x in self.COLS]; ws.append(headers)
        header_fill = PatternFill("solid", fgColor="D9EAF7")
        warn_fill = PatternFill("solid", fgColor="FFF2CC")
        for cell in ws[1]:
            cell.font = Font(bold=True); cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in self.rows:
            ws.append([getattr(row, key) for _, key in self.COLS])
            excel_row = ws.max_row
            if row.note or row.confidence < 0.72:
                for cell in ws[excel_row]: cell.fill = warn_fill
            for cell in ws[excel_row]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        widths = [8,9,14,18,16,24,27,55,15,15,16,46,12,12]
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        wb.save(out)
        QMessageBox.information(self, "저장 완료", f"Excel 파일을 저장했습니다.\n{out}")


def main():
    app = QApplication(sys.argv)
    app.setApplicationName(APP_NAME)
    win = MainWindow(); win.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
